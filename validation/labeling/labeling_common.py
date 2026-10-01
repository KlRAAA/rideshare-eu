"""
Shared pieces of the human-labeling kit: turning a dataset row into what an
evaluator sees, the workbook column layout, and reading labels back out.

Evaluators see the situation (times, seats, genders, preferences, distances,
a map link) and never the algorithm's output (route overlap, schedule
alignment, preference match, score, Stage 1 result) — otherwise the "ground
truth" would just echo PSGA back at itself.
"""

import json
import math
import os

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
VALIDATION_DIR = os.path.dirname(HERE)
DATASET_PATH = os.path.join(VALIDATION_DIR, "dataset_500_pairs.json")

ACCEPTABLE = "ACCEPTABLE"
NOT_ACCEPTABLE = "NOT ACCEPTABLE"
LABELS = (ACCEPTABLE, NOT_ACCEPTABLE)

SHEET_NAME = "Pairs"

# (header, width) in sheet order. The last two are the only editable columns.
COLUMNS = [
    ("Pair ID", 11),
    ("Passenger leaves at", 12),
    ("Passenger can wait (± min)", 12),
    ("Passenger gender", 10),
    ("Passenger wants a same-gender driver", 14),
    ("Passenger already knows this driver", 14),
    ("Driver leaves at", 12),
    ("Gap between departures (min)", 13),
    ("Seats left", 8),
    ("Driver gender", 10),
    ("Driver accepts", 16),
    ("Driver takes only riders they know", 14),
    ("Passenger home → campus (km)", 13),
    ("Driver start → passenger home (km)", 14),
    ("Extra distance for the driver to pick up (km)", 15),
    ("Map: driver start → passenger → campus", 16),
    ("Judgment", 18),
    ("Notes (optional)", 40),
]
HEADERS = [h for h, _ in COLUMNS]
JUDGMENT_COL = HEADERS.index("Judgment") + 1  # 1-based for openpyxl
NOTES_COL = HEADERS.index("Notes (optional)") + 1
PAIR_ID_COL = 1


def load_dataset(path=DATASET_PATH):
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def haversine_km(a, b):
    lat1, lng1, lat2, lng2 = map(math.radians, (a[0], a[1], b[0], b[1]))
    h = math.sin((lat2 - lat1) / 2) ** 2 + math.cos(lat1) * math.cos(lat2) * math.sin((lng2 - lng1) / 2) ** 2
    return 2 * 6371.0 * math.asin(math.sqrt(h))


def clock(minutes):
    """Minutes after midnight -> '7:05 AM'."""
    hours, mins = divmod(int(minutes) % (24 * 60), 60)
    suffix = "AM" if hours < 12 else "PM"
    return f"{(hours % 12) or 12}:{mins:02d} {suffix}"


def yes_no(value):
    return "Yes" if value else "No"


def map_link(driver_start, passenger_home, campus):
    points = "/".join(f"{lat:.6f},{lng:.6f}" for lat, lng in (driver_start, passenger_home, campus))
    return f"https://www.google.com/maps/dir/{points}"


def evaluator_row(r):
    """What an evaluator sees for one pair — inputs only, no algorithm output."""
    passenger_home = r["passengerOrigin"]
    driver_start = r["tripOrigin"]
    campus = r["passengerDestination"]

    home_to_campus = haversine_km(passenger_home, campus)
    driver_to_home = haversine_km(driver_start, passenger_home)
    driver_to_campus = haversine_km(driver_start, campus)
    # Straight-line detour: how much longer the driver's trip gets by passing
    # the passenger's home on the way to campus.
    extra = max(0.0, driver_to_home + home_to_campus - driver_to_campus)

    return [
        r["tripId"],
        clock(r["passengerDepartureMinutes"]),
        r["passengerFlexWindowMinutes"],
        r["passengerGender"].title(),
        yes_no(r["passengerSameGenderOnly"]),
        yes_no(r["passengerIsFamiliarWithHost"]),
        clock(r["tripDepartureMinutes"]),
        abs(r["tripDepartureMinutes"] - r["passengerDepartureMinutes"]),
        r["tripSeatsAvailable"],
        r["hostGender"].title(),
        "Same gender only" if r["hostGenderPreference"] == "SAME_GENDER" else "Anyone",
        yes_no(r["familiarRidersOnly"]),
        round(home_to_campus, 2),
        round(driver_to_home, 2),
        round(extra, 2),
        map_link(driver_start, passenger_home, campus),
    ]


def normalize_label(value):
    """Returns ACCEPTABLE / NOT ACCEPTABLE, None for an empty cell, or raises ValueError."""
    if value is None:
        return None
    text = " ".join(str(value).upper().replace("_", " ").split())
    if text == "":
        return None
    if text in LABELS:
        return text
    raise ValueError(f"not a valid judgment: {value!r}")


def read_labels(path):
    """{pair_id: label or None} plus a list of (pair_id, bad value) for unreadable cells."""
    wb = load_workbook(path, read_only=True, data_only=True)
    labels, invalid = {}, []
    try:
        for row in wb[SHEET_NAME].iter_rows(min_row=2, values_only=True):
            pair_id = row[PAIR_ID_COL - 1]
            if not pair_id:
                continue
            raw = row[JUDGMENT_COL - 1] if len(row) >= JUDGMENT_COL else None
            try:
                labels[pair_id] = normalize_label(raw)
            except ValueError:
                labels[pair_id] = None
                invalid.append((pair_id, raw))
    finally:
        wb.close()  # read-only mode keeps the file open on Windows until closed
    return labels, invalid
