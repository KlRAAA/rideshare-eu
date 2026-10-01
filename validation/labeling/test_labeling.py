"""
Tests for the labeling kit. Run from the repo root:

    python -m unittest discover -s validation/labeling -v

Every label written here is test data inside a temporary folder — nothing in
this file produces or touches the real ground_truth.json.
"""

import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest

from openpyxl import load_workbook

from build_labeling_kit import EVALUATOR_SEEDS, build_workbook, shuffled_rows
from labeling_common import (
    ACCEPTABLE, HEADERS, JUDGMENT_COL, NOT_ACCEPTABLE, SHEET_NAME, VALIDATION_DIR,
    clock, evaluator_row, load_dataset, normalize_label, read_labels,
)
import merge_labels

DATASET = load_dataset()
SMALL = DATASET[:6]  # one query's worth (5 pairs) plus one from the next


def fill(path, labels):
    """Writes {pair_id: label} into a workbook's Judgment column, like an evaluator would."""
    wb = load_workbook(path)
    ws = wb[SHEET_NAME]
    for row in range(2, ws.max_row + 1):
        pid = ws.cell(row=row, column=1).value
        if pid in labels:
            ws.cell(row=row, column=JUDGMENT_COL, value=labels[pid])
    wb.save(path)


class EvaluatorViewTests(unittest.TestCase):
    def test_clock_formats_minutes_after_midnight(self):
        self.assertEqual(clock(0), "12:00 AM")
        self.assertEqual(clock(452), "7:32 AM")
        self.assertEqual(clock(720), "12:00 PM")
        self.assertEqual(clock(990), "4:30 PM")

    def test_row_matches_the_headers_minus_the_two_editable_columns(self):
        row = evaluator_row(DATASET[0])
        self.assertEqual(len(row), len(HEADERS) - 2)
        self.assertEqual(row[0], DATASET[0]["tripId"])

    def test_evaluators_never_see_the_algorithm_output(self):
        for header in HEADERS:
            for word in ("score", "overlap", "alignment", "preference match", "stage", "psga"):
                self.assertNotIn(word, header.lower())

    def test_row_shows_the_inputs_in_plain_words(self):
        r = dict(DATASET[0], hostGenderPreference="SAME_GENDER", familiarRidersOnly=True,
                 passengerDepartureMinutes=420, tripDepartureMinutes=432)
        row = dict(zip(HEADERS, evaluator_row(r)))
        self.assertEqual(row["Passenger leaves at"], "7:00 AM")
        self.assertEqual(row["Driver leaves at"], "7:12 AM")
        self.assertEqual(row["Gap between departures (min)"], 12)
        self.assertEqual(row["Driver accepts"], "Same gender only")
        self.assertEqual(row["Driver takes only riders they know"], "Yes")
        self.assertGreaterEqual(row["Extra distance for the driver to pick up (km)"], 0)
        self.assertTrue(row["Map: driver start → passenger → campus"].startswith("https://www.google.com/maps/dir/"))

    def test_each_evaluator_gets_all_500_pairs_once_in_a_different_fixed_order(self):
        a = [r[0] for r in shuffled_rows(DATASET, EVALUATOR_SEEDS["A"])]
        b = [r[0] for r in shuffled_rows(DATASET, EVALUATOR_SEEDS["B"])]
        ids = sorted(r["tripId"] for r in DATASET)
        self.assertEqual(sorted(a), ids)
        self.assertEqual(sorted(b), ids)
        self.assertNotEqual(a, b)
        self.assertEqual(a, [r[0] for r in shuffled_rows(DATASET, EVALUATOR_SEEDS["A"])])


class LabelReadingTests(unittest.TestCase):
    def test_normalize_label_accepts_case_and_spacing_variants(self):
        self.assertEqual(normalize_label(" acceptable "), ACCEPTABLE)
        self.assertEqual(normalize_label("Not  Acceptable"), NOT_ACCEPTABLE)
        self.assertEqual(normalize_label("NOT_ACCEPTABLE"), NOT_ACCEPTABLE)
        self.assertIsNone(normalize_label(None))
        self.assertIsNone(normalize_label("   "))

    def test_normalize_label_rejects_anything_else(self):
        for bad in ("yes", "maybe", "1", "ACCEPT"):
            with self.assertRaises(ValueError):
                normalize_label(bad)

    def test_workbook_round_trip(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "Evaluator_A.xlsx")
            build_workbook([evaluator_row(r) for r in SMALL], "A", path)
            labels, invalid = read_labels(path)
            self.assertEqual(set(labels), {r["tripId"] for r in SMALL})
            self.assertTrue(all(v is None for v in labels.values()))

            ids = [r["tripId"] for r in SMALL]
            fill(path, {ids[0]: ACCEPTABLE, ids[1]: "not acceptable", ids[2]: "dunno"})
            labels, invalid = read_labels(path)
            self.assertEqual(labels[ids[0]], ACCEPTABLE)
            self.assertEqual(labels[ids[1]], NOT_ACCEPTABLE)
            self.assertIsNone(labels[ids[2]])
            self.assertEqual(invalid, [(ids[2], "dunno")])

    def test_only_judgment_and_notes_are_editable(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "Evaluator_A.xlsx")
            build_workbook([evaluator_row(r) for r in SMALL], "A", path)
            ws = load_workbook(path)[SHEET_NAME]
            self.assertTrue(ws.protection.sheet)
            unlocked = {c.column for c in ws[2] if not c.protection.locked}
            self.assertEqual(unlocked, {JUDGMENT_COL, JUDGMENT_COL + 1})


class KappaTests(unittest.TestCase):
    def test_textbook_example(self):
        a = {"1": ACCEPTABLE, "2": ACCEPTABLE, "3": NOT_ACCEPTABLE, "4": NOT_ACCEPTABLE}
        b = {"1": ACCEPTABLE, "2": NOT_ACCEPTABLE, "3": NOT_ACCEPTABLE, "4": NOT_ACCEPTABLE}
        # observed 0.75, chance 0.5 -> kappa 0.5
        self.assertAlmostEqual(merge_labels.cohens_kappa(a, b), 0.5)

    def test_perfect_agreement_is_one_even_when_everyone_says_the_same(self):
        same = {"1": ACCEPTABLE, "2": ACCEPTABLE}
        self.assertEqual(merge_labels.cohens_kappa(same, dict(same)), 1.0)

    def test_reading_bands(self):
        self.assertEqual(merge_labels.kappa_reading(0.85), "almost perfect")
        self.assertEqual(merge_labels.kappa_reading(0.65), "substantial")
        self.assertEqual(merge_labels.kappa_reading(0.30), "fair")


class MergeTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.truth_path = os.path.join(self.tmp, "ground_truth.json")
        self.ids = [r["tripId"] for r in SMALL]
        for who in ("A", "B"):
            build_workbook([evaluator_row(r) for r in SMALL], who, os.path.join(self.tmp, f"Evaluator_{who}.xlsx"))

    def tearDown(self):
        shutil.rmtree(self.tmp)

    def run_merge(self):
        return merge_labels.run(work_dir=self.tmp, ground_truth_path=self.truth_path, dataset=SMALL)

    def label_all(self, who, overrides=None):
        labels = {pid: ACCEPTABLE if i % 2 == 0 else NOT_ACCEPTABLE for i, pid in enumerate(self.ids)}
        labels.update(overrides or {})
        fill(os.path.join(self.tmp, f"Evaluator_{who}.xlsx"), labels)
        return labels

    def test_refuses_to_merge_an_unfinished_sheet(self):
        self.label_all("A")
        fill(os.path.join(self.tmp, "Evaluator_B.xlsx"), {self.ids[0]: ACCEPTABLE})
        with self.assertRaises(merge_labels.LabelingIncomplete) as ctx:
            self.run_merge()
        self.assertIn("Evaluator_B.xlsx", str(ctx.exception))
        self.assertIn("no judgment yet", str(ctx.exception))
        self.assertFalse(os.path.exists(self.truth_path))

    def test_full_agreement_writes_ground_truth_directly(self):
        labels = self.label_all("A")
        self.label_all("B")
        truth = self.run_merge()
        self.assertEqual(truth, {pid: label == ACCEPTABLE for pid, label in labels.items()})
        with open(self.truth_path) as f:
            self.assertEqual(json.load(f), truth)
        self.assertFalse(os.path.exists(os.path.join(self.tmp, "Tiebreaker.xlsx")))

    def test_disagreements_go_to_the_tiebreaker_who_decides_them(self):
        a = self.label_all("A")
        flipped = self.ids[1]  # A says NOT ACCEPTABLE
        self.label_all("B", {flipped: ACCEPTABLE})

        self.assertIsNone(self.run_merge())
        self.assertFalse(os.path.exists(self.truth_path))
        tiebreaker = os.path.join(self.tmp, "Tiebreaker.xlsx")
        pending, _ = read_labels(tiebreaker)
        self.assertEqual(list(pending), [flipped])

        with open(os.path.join(self.tmp, "agreement_report.json")) as f:
            report = json.load(f)
        self.assertEqual(report["disagreed"], 1)
        self.assertEqual(report["disagreement_ids"], [flipped])

        fill(tiebreaker, {flipped: ACCEPTABLE})
        truth = self.run_merge()
        self.assertTrue(truth[flipped])
        for pid in self.ids:
            if pid != flipped:
                self.assertEqual(truth[pid], a[pid] == ACCEPTABLE)

    def test_an_unfinished_tiebreaker_sheet_blocks_the_merge_and_is_never_overwritten(self):
        self.label_all("A")
        self.label_all("B", {self.ids[1]: ACCEPTABLE})
        self.run_merge()
        tiebreaker = os.path.join(self.tmp, "Tiebreaker.xlsx")
        stamp = os.path.getmtime(tiebreaker)
        with self.assertRaises(merge_labels.LabelingIncomplete):
            self.run_merge()
        self.assertEqual(os.path.getmtime(tiebreaker), stamp)


class ScorerTests(unittest.TestCase):
    def test_scorer_reads_merged_ground_truth_and_saves_results(self):
        """If every query's PSGA #1 pick is the only acceptable pair, PSGA scores precision@1 = 1."""
        with tempfile.TemporaryDirectory() as tmp:
            for name in ("score_against_ground_truth.py", "dataset_500_pairs.json", "method_rankings.json"):
                shutil.copy(os.path.join(VALIDATION_DIR, name), tmp)
            with open(os.path.join(tmp, "method_rankings.json")) as f:
                rankings = json.load(f)
            top_picks = {ranked["psga"][0] for ranked in rankings.values() if ranked["psga"]}
            truth = {r["tripId"]: r["tripId"] in top_picks for r in DATASET}
            with open(os.path.join(tmp, "ground_truth.json"), "w") as f:
                json.dump(truth, f)

            subprocess.run([sys.executable, "score_against_ground_truth.py"], cwd=tmp, check=True, capture_output=True)

            with open(os.path.join(tmp, "precision_results.json")) as f:
                results = json.load(f)
            self.assertEqual(results["labeled_pairs"], 500)
            self.assertEqual(results["methods"]["psga"]["precision@1"], 1.0)
            self.assertEqual(results["methods"]["psga"]["mean_rank_of_first_hit"], 1.0)


if __name__ == "__main__":
    unittest.main()
