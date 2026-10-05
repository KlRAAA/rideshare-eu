"""
Builds the blank labeling workbooks for the PSGA ground truth (manuscript,
Algorithm Validation): one per evaluator, each holding all 500 pairs of
dataset_500_pairs.json in its own shuffled order.

    python validation/labeling/build_labeling_kit.py

Writes Evaluator_A.xlsx and Evaluator_B.xlsx next to this script. Refuses to
overwrite a workbook that already exists, so a half-labeled sheet can't be
wiped by re-running it; pass --force to rebuild blank ones on purpose.

The tiebreaker's workbook is not built here — merge_labels.py builds it from
the pairs A and B disagree on, once both are done.
"""

import os
import random
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation

from labeling_common import (
    ACCEPTABLE, COLUMNS, HERE, JUDGMENT_COL, NOTES_COL, NOT_ACCEPTABLE,
    SHEET_NAME, evaluator_row, load_dataset,
)

# Fixed seeds: anyone re-running this gets the same row orders.
EVALUATOR_SEEDS = {"A": 20261001, "B": 20261002}

RUBRIC = [
    ("How to judge each pair", None),
    ("Question: would a reasonable MSEUF student accept being offered this ride, and would this driver reasonably take them?", None),
    ("Mark the pair NOT ACCEPTABLE if any of these is true:", None),
    ("1.", "Seats left is 0."),
    ("2.", "Passenger wants Women+ trips only and the driver accepts anyone."),
    ("3.", "Driver accepts Women+ only (women and non-binary riders) and the passenger is a man or prefers not to say."),
    ("4.", "Driver takes only riders they know and the passenger does not know this driver."),
    ("Otherwise, judge time and route:", None),
    ("Time", "Gap within the passenger's ± wait → fine. Up to 10 minutes beyond it → your judgment. More than that → NOT ACCEPTABLE."),
    ("Route", "Extra distance up to 1 km → fine. 1 to 3 km → your judgment (open the map link: is the pickup roughly on the way?). More than 3 km → NOT ACCEPTABLE."),
    ("Result", "ACCEPTABLE only if rules 1–4 pass and both time and route are fine or judged fine."),
    ("Distances are straight-line kilometres. The map link shows the actual roads.", None),
]

INSTRUCTIONS = [
    "PSGA ground-truth labeling — Evaluator {who}",
    "",
    "Each row is one passenger and one driver's trip to MSEUF. Decide whether it is an acceptable match,",
    "using the Rubric tab. Pick ACCEPTABLE or NOT ACCEPTABLE in the Judgment column for every row.",
    "",
    "Rules for a fair result:",
    "  • Work alone. Don't discuss pairs with the other evaluator until you have both finished.",
    "  • Use only the facts in the row and the map link. The algorithm's own scores are deliberately left out.",
    "  • Don't change the rubric partway through. If a case doesn't fit it, decide and write why in Notes.",
    "  • Only the Judgment and Notes columns can be edited; the rest of the sheet is locked so facts can't be changed by accident.",
    "",
    "When every row is filled in, save the file and send it back unchanged in name (Evaluator_{who}.xlsx).",
]

HEADER_FILL = PatternFill("solid", fgColor="DDE7F0")
EDIT_FILL = PatternFill("solid", fgColor="FFF7D6")


def write_text_sheet(ws, lines, width=110):
    ws.column_dimensions["A"].width = width
    for i, line in enumerate(lines, start=1):
        ws.cell(row=i, column=1, value=line)
    ws["A1"].font = Font(bold=True, size=14)


def write_rubric(ws):
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 100
    for i, (left, right) in enumerate(RUBRIC, start=1):
        ws.cell(row=i, column=1, value=left)
        if right:
            ws.cell(row=i, column=2, value=right).alignment = Alignment(wrap_text=True, vertical="top")
        else:
            ws.cell(row=i, column=1).font = Font(bold=True)
    ws["A1"].font = Font(bold=True, size=14)


def write_pairs_sheet(ws, rows):
    for col, (header, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[cell.column_letter].width = width
    ws.row_dimensions[1].height = 62
    ws.freeze_panes = "B2"

    map_col = len(COLUMNS) - 2  # the map link sits just before Judgment
    for r, values in enumerate(rows, start=2):
        for c, value in enumerate(values, start=1):
            ws.cell(row=r, column=c, value=value)
        link = ws.cell(row=r, column=map_col)
        link.hyperlink = values[map_col - 1]
        link.value = "Open map"
        link.font = Font(color="1F5FA8", underline="single")
        for col in (JUDGMENT_COL, NOTES_COL):
            cell = ws.cell(row=r, column=col)
            cell.protection = Protection(locked=False)
            cell.fill = EDIT_FILL

    last = len(rows) + 1
    judgment = DataValidation(
        type="list",
        formula1=f'"{ACCEPTABLE},{NOT_ACCEPTABLE}"',
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="Pick a judgment",
        error=f"Choose {ACCEPTABLE} or {NOT_ACCEPTABLE}.",
    )
    letter = ws.cell(row=1, column=JUDGMENT_COL).column_letter
    judgment.add(f"{letter}2:{letter}{last}")
    ws.add_data_validation(judgment)

    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(COLUMNS)).column_letter}{last}"
    ws.protection.sheet = True
    ws.protection.autoFilter = False  # filtering stays allowed on the locked sheet


def build_workbook(rows, who, path):
    wb = Workbook()
    write_text_sheet(wb.active, [line.format(who=who) for line in INSTRUCTIONS])
    wb.active.title = "Instructions"
    write_rubric(wb.create_sheet("Rubric"))
    write_pairs_sheet(wb.create_sheet(SHEET_NAME), rows)
    wb.active = 2
    wb.save(path)


def shuffled_rows(dataset, seed):
    rows = [evaluator_row(r) for r in dataset]
    random.Random(seed).shuffle(rows)
    return rows


def main(argv):
    force = "--force" in argv
    dataset = load_dataset()
    for who, seed in EVALUATOR_SEEDS.items():
        path = os.path.join(HERE, f"Evaluator_{who}.xlsx")
        if os.path.exists(path) and not force:
            print(f"{os.path.basename(path)} already exists, left untouched (use --force to rebuild a blank one).")
            continue
        build_workbook(shuffled_rows(dataset, seed), who, path)
        print(f"Wrote {os.path.basename(path)}: {len(dataset)} pairs, judgments blank.")


if __name__ == "__main__":
    main(sys.argv[1:])
